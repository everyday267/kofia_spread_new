"""data/spread_history.csv → 잔존만기별 탭 엑셀 (메일 첨부 엑셀과 같은 양식).

각 탭: 기준일(YYYY-MM-DD) | 국고채금리 | AAA | AA+ | ... | BBB-   (기준일 내림차순)
값 단위: 국고채금리 = %, 등급 = 국고 대비 스프레드 %p  (KOFIA 원값 그대로)

    python spread_export.py [출력경로]     # 기본: docs/kofia_corp_spread_history.xlsx
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import spread_analytics as sa

EXPORT_PATH = Path(os.getenv("EXPORT_PATH", "docs/kofia_corp_spread_history.xlsx"))
CORP_GRADES = ["AAA", "AA+", "AA", "AA-", "A+", "A", "A-", "BBB+", "BBB", "BBB-"]


def tenor_sheet_name(code: str) -> str:
    m = int(code)
    y, r = divmod(m, 12)
    return (f"{m}개월" if y == 0 else f"{y}년" + (f"{r}개월" if r else "")) + "물"


def tenor_frame(hist: pd.DataFrame, code: str) -> pd.DataFrame:
    """메일 엑셀과 같은 열 구성의 wide 표."""
    h = hist[hist["tenor"] == str(code)]
    w = h.pivot_table(index="date", columns="grade", values="value", aggfunc="last")
    grades = [g for g in CORP_GRADES if g in w.columns] + \
             [g for g in sa.GRADE_ORDER if g in w.columns and g not in CORP_GRADES]
    out = w[[sa.GOV] + grades].rename(columns={sa.GOV: "국고채금리"})
    out = out.sort_index(ascending=False)
    out.index = [f"{d[:4]}-{d[4:6]}-{d[6:]}" for d in out.index]
    out.index.name = "기준일"
    return out.reset_index()


def build_workbook(hist: pd.DataFrame, out: Path = EXPORT_PATH) -> Path:
    codes = sorted(hist["tenor"].dropna().unique(), key=int)
    out.parent.mkdir(parents=True, exist_ok=True)
    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    head_fill = PatternFill("solid", fgColor="F2F2F2")
    thin = Side(style="thin", color="BFBFBF")
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for code in codes:
            df = tenor_frame(hist, code)
            name = tenor_sheet_name(code)
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
            for c, col in enumerate(df.columns, start=1):
                ws.column_dimensions[get_column_letter(c)].width = 12 if c == 1 else 10
                h = ws.cell(row=1, column=c)
                h.font, h.fill, h.border = bold, head_fill, Border(bottom=thin)
                h.alignment = Alignment(horizontal="center")
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = font
                    if cell.column == 1:
                        cell.alignment = Alignment(horizontal="center")
                    else:
                        cell.number_format = "0.00"
    return out


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else EXPORT_PATH
    print(build_workbook(sa.load_history(), target))
